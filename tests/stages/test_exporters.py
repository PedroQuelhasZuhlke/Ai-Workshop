"""Stage tests for `export` (`ExportStage`): golden-path files for every supplier and
format, one test per validation-gate failure, formula-injection neutralization in
both CSV and XLSX bytes, a CSV round-trip, and a determinism check, per
`.github/instructions/pipeline-stages.instructions.md` and
`.github/skills/supplier-export/SKILL.md`.
"""

import csv
import io
from datetime import UTC, datetime
from decimal import Decimal
from random import Random

import pytest
from openpyxl import load_workbook

from pcbpick.contracts.enums import (
    DecisionMode,
    ExportFormat,
    PlanWarningKind,
    SubstitutionTier,
    SupplierId,
)
from pcbpick.contracts.plan import (
    PlanLine,
    PlanTotals,
    PlanWarning,
    PurchasePlan,
    UnfilledRequirement,
)
from pcbpick.exporters.stage import ExportStage, ExportStageInput
from pcbpick.pipeline.context import PipelineConfig, StageContext
from pcbpick.pipeline.results import Severity


def _make_ctx() -> StageContext:
    return StageContext(
        run_id="test-run",
        now=datetime(2026, 9, 22, tzinfo=UTC),
        rng=Random("test-run"),
        config=PipelineConfig(),
    )


def _line(
    requirement_id: str,
    supplier: SupplierId,
    *,
    mpn: str = "RC0603FR-0710KL",
    sku: str = "RC0603-SKU-1",
    substitution_tier: SubstitutionTier = SubstitutionTier.TIER_0,
    order_qty: int = 100,
    snapshot_id: str = "SNAP-1",
) -> PlanLine:
    return PlanLine(
        requirement_id=requirement_id,
        mpn=mpn,
        manufacturer="Yageo",
        supplier=supplier,
        sku=sku,
        substitution_tier=substitution_tier,
        order_qty=order_qty,
        unit_price=Decimal("0.05"),
        extended_price=Decimal("0.05") * order_qty,
        currency="USD",
        snapshot_id=snapshot_id,
    )


def _make_plan(
    lines: list[PlanLine],
    unfilled: list[UnfilledRequirement] | None = None,
    warnings: list[PlanWarning] | None = None,
) -> PurchasePlan:
    grand_total = sum((line.extended_price for line in lines), Decimal("0"))
    return PurchasePlan(
        mode=DecisionMode.LOWEST_COST,
        board_quantity=1,
        attrition_pct=Decimal("0"),
        lines=lines,
        totals=PlanTotals(
            currency="USD",
            supplier_subtotals={},
            shipping_total=Decimal("0"),
            grand_total=grand_total,
        ),
        longest_lead_time_days=0,
        supplier_count=len({line.supplier for line in lines}),
        unfilled=unfilled or [],
        warnings=warnings or [],
        snapshot_ids=sorted({line.snapshot_id for line in lines}),
    )


def _base_input(
    plan: PurchasePlan, supplier: SupplierId, export_format: ExportFormat, **overrides: object
) -> ExportStageInput:
    defaults: dict[str, object] = {
        "plan": plan,
        "supplier": supplier,
        "format": export_format,
        "approved_by": "alice",
        "approved_at": datetime(2026, 9, 22, tzinfo=UTC),
        "stale_offer_snapshot_ids": frozenset(),
        "reviewed_high_tier_requirement_ids": frozenset(),
    }
    defaults.update(overrides)
    return ExportStageInput(**defaults)  # type: ignore[arg-type]


@pytest.mark.parametrize(
    "supplier",
    [SupplierId.MOUSER, SupplierId.DIGIKEY, SupplierId.FARNELL],
)
@pytest.mark.parametrize("export_format", [ExportFormat.CSV, ExportFormat.XLSX])
def test_golden_path_produces_export_and_file_bytes(
    supplier: SupplierId, export_format: ExportFormat
) -> None:
    plan = _make_plan([_line("REQ-R1", supplier)])
    result = ExportStage().run(_base_input(plan, supplier, export_format), _make_ctx())

    assert result.output is not None
    export = result.output.export
    assert export.supplier == supplier
    assert export.format == export_format
    assert export.approved_by == "alice"
    assert len(export.rows) == 1
    assert result.output.file_bytes
    assert not any(issue.severity == Severity.BLOCKING for issue in result.issues)


def test_mouser_row_maps_expected_columns() -> None:
    plan = _make_plan(
        [_line("REQ-R1", SupplierId.MOUSER, sku="MOUSER-SKU-1", mpn="RC0603FR-0710KL")]
    )
    result = ExportStage().run(_base_input(plan, SupplierId.MOUSER, ExportFormat.CSV), _make_ctx())
    assert result.output is not None
    row = result.output.export.rows[0]
    assert row == {
        "Mouser Part Number": "MOUSER-SKU-1",
        "Manufacturer Part Number": "RC0603FR-0710KL",
        "Quantity": "100",
        "Customer Part Number": "REQ-R1",
    }


def test_digikey_row_maps_expected_columns() -> None:
    plan = _make_plan([_line("REQ-R1", SupplierId.DIGIKEY, sku="RC0603-1-ND")])
    result = ExportStage().run(_base_input(plan, SupplierId.DIGIKEY, ExportFormat.CSV), _make_ctx())
    assert result.output is not None
    row = result.output.export.rows[0]
    assert row == {
        "Quantity": "100",
        "Digi-Key Part Number": "RC0603-1-ND",
        "Manufacturer Part Number": "RC0603FR-0710KL",
        "Customer Reference": "REQ-R1",
    }


def test_farnell_row_maps_expected_columns() -> None:
    plan = _make_plan([_line("REQ-R1", SupplierId.FARNELL, sku="FARNELL-9999")])
    result = ExportStage().run(_base_input(plan, SupplierId.FARNELL, ExportFormat.CSV), _make_ctx())
    assert result.output is not None
    row = result.output.export.rows[0]
    assert row == {
        "Order Code": "FARNELL-9999",
        "Quantity": "100",
        "Line Note": "REQ-R1",
    }


def test_export_only_includes_lines_for_the_requested_supplier() -> None:
    plan = _make_plan(
        [
            _line("REQ-R1", SupplierId.MOUSER, sku="MOUSER-SKU-1"),
            _line("REQ-C1", SupplierId.DIGIKEY, sku="DIGIKEY-SKU-1"),
        ]
    )
    result = ExportStage().run(_base_input(plan, SupplierId.MOUSER, ExportFormat.CSV), _make_ctx())
    assert result.output is not None
    assert len(result.output.export.rows) == 1
    assert result.output.export.rows[0]["Customer Part Number"] == "REQ-R1"


def test_export_plan_hash_matches_compute_plan_hash() -> None:
    from pcbpick.planning.plan_hash import compute_plan_hash

    plan = _make_plan([_line("REQ-R1", SupplierId.MOUSER)])
    result = ExportStage().run(_base_input(plan, SupplierId.MOUSER, ExportFormat.CSV), _make_ctx())
    assert result.output is not None
    assert result.output.export.plan_hash == compute_plan_hash(plan)


# --- Validation gate failures ------------------------------------------------------


def test_unapproved_plan_is_rejected() -> None:
    plan = _make_plan([_line("REQ-R1", SupplierId.MOUSER)])
    result = ExportStage().run(
        _base_input(plan, SupplierId.MOUSER, ExportFormat.CSV, approved_by=""), _make_ctx()
    )
    assert result.output is None
    blocking = [i for i in result.issues if i.severity == Severity.BLOCKING]
    assert blocking
    assert any("approv" in i.message for i in blocking)


def test_unacknowledged_stale_offer_warning_is_rejected() -> None:
    plan = _make_plan(
        [_line("REQ-R1", SupplierId.MOUSER)],
        warnings=[
            PlanWarning(
                kind=PlanWarningKind.STALE_OFFER,
                requirement_id="REQ-R1",
                message="offer fetched 200 days ago",
            )
        ],
    )
    result = ExportStage().run(_base_input(plan, SupplierId.MOUSER, ExportFormat.CSV), _make_ctx())
    assert result.output is None
    blocking = [i for i in result.issues if i.severity == Severity.BLOCKING]
    assert blocking
    assert any("stale" in i.message for i in blocking)


def test_acknowledged_stale_offer_warning_allows_export_and_is_recorded() -> None:
    """An engineer may override a *specific* stale line by acknowledging its
    requirement at approval time (`stale_offer_snapshot_ids`, despite its name --
    see `exporters.validation.check_no_stale_offers`'s docstring). The override
    must still be visible in the run's own issues, not just the approval ledger."""
    plan = _make_plan(
        [_line("REQ-R1", SupplierId.MOUSER)],
        warnings=[
            PlanWarning(
                kind=PlanWarningKind.STALE_OFFER,
                requirement_id="REQ-R1",
                message="offer fetched 200 days ago",
            )
        ],
    )
    result = ExportStage().run(
        _base_input(
            plan,
            SupplierId.MOUSER,
            ExportFormat.CSV,
            stale_offer_snapshot_ids=frozenset({"REQ-R1"}),
        ),
        _make_ctx(),
    )
    assert result.output is not None
    assert not any(i.severity == Severity.BLOCKING for i in result.issues)
    override_issues = [
        i for i in result.issues if i.code == "exported_with_acknowledged_stale_offer"
    ]
    assert len(override_issues) == 1
    assert override_issues[0].item_id == "REQ-R1"


def test_plans_own_stale_offer_warning_is_rejected_even_without_caller_flagging_it() -> None:
    """The gate must not fail open: a `STALE_OFFER` warning already on the plan
    (as `PlanningStage` would emit) is enough on its own -- `stale_offer_snapshot_ids`
    is left at its empty default here, exactly as a careless caller might do."""
    plan = _make_plan(
        [_line("REQ-R1", SupplierId.MOUSER, snapshot_id="SNAP-1")],
        warnings=[
            PlanWarning(
                kind=PlanWarningKind.STALE_OFFER,
                requirement_id="REQ-R1",
                message="offer fetched 200 days ago",
            )
        ],
    )
    result = ExportStage().run(_base_input(plan, SupplierId.MOUSER, ExportFormat.CSV), _make_ctx())
    assert result.output is None
    blocking = [i for i in result.issues if i.severity == Severity.BLOCKING]
    assert blocking
    assert any("stale" in i.message for i in blocking)


def test_unreviewed_tier2_line_is_rejected() -> None:
    plan = _make_plan(
        [_line("REQ-R1", SupplierId.MOUSER, substitution_tier=SubstitutionTier.TIER_2)]
    )
    result = ExportStage().run(_base_input(plan, SupplierId.MOUSER, ExportFormat.CSV), _make_ctx())
    assert result.output is None
    blocking = [i for i in result.issues if i.severity == Severity.BLOCKING]
    assert blocking
    assert any("tier" in i.message for i in blocking)


def test_reviewed_tier2_line_passes() -> None:
    plan = _make_plan(
        [_line("REQ-R1", SupplierId.MOUSER, substitution_tier=SubstitutionTier.TIER_2)]
    )
    result = ExportStage().run(
        _base_input(
            plan,
            SupplierId.MOUSER,
            ExportFormat.CSV,
            reviewed_high_tier_requirement_ids=frozenset({"REQ-R1"}),
        ),
        _make_ctx(),
    )
    assert result.output is not None


def test_unfilled_requirements_reject_export() -> None:
    plan = _make_plan(
        [_line("REQ-R1", SupplierId.MOUSER)],
        unfilled=[UnfilledRequirement(requirement_id="REQ-C1", reason="no viable offer")],
    )
    result = ExportStage().run(_base_input(plan, SupplierId.MOUSER, ExportFormat.CSV), _make_ctx())
    assert result.output is None
    blocking = [i for i in result.issues if i.severity == Severity.BLOCKING]
    assert blocking
    assert any("unfilled" in i.message for i in blocking)


# --- Formula injection --------------------------------------------------------------


@pytest.mark.parametrize("risky_prefix", ["=", "+", "-", "@", "\t", "\r"])
def test_formula_injection_neutralized_in_csv(risky_prefix: str) -> None:
    sku = f"{risky_prefix}cmd|calc"
    plan = _make_plan([_line("REQ-R1", SupplierId.MOUSER, sku=sku)])
    result = ExportStage().run(_base_input(plan, SupplierId.MOUSER, ExportFormat.CSV), _make_ctx())
    assert result.output is not None
    text = result.output.file_bytes.decode("utf-8-sig")
    rows = list(csv.reader(io.StringIO(text)))
    header, data_row = rows[0], rows[1]
    part_number_cell = data_row[header.index("Mouser Part Number")]
    assert part_number_cell.startswith("'")
    assert part_number_cell[1:] == sku


@pytest.mark.parametrize("risky_prefix", ["=", "+", "-", "@", "\t", "\r"])
def test_formula_injection_neutralized_in_xlsx(risky_prefix: str) -> None:
    sku = f"{risky_prefix}cmd|calc"
    plan = _make_plan([_line("REQ-R1", SupplierId.MOUSER, sku=sku)])
    result = ExportStage().run(_base_input(plan, SupplierId.MOUSER, ExportFormat.XLSX), _make_ctx())
    assert result.output is not None
    workbook = load_workbook(io.BytesIO(result.output.file_bytes))
    sheet = workbook.active
    assert sheet is not None
    header = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    data_row = [cell.value for cell in next(sheet.iter_rows(min_row=2, max_row=2))]
    part_number_cell = data_row[header.index("Mouser Part Number")]
    assert isinstance(part_number_cell, str)
    assert part_number_cell.startswith("'")
    # OOXML is XML under the hood, and XML parsers normalize a bare "\r" to "\n" per
    # the XML spec's end-of-line handling (section 2.11) -- inherent to the file
    # format, not a neutralization bug, so the readback comparison accounts for it.
    assert part_number_cell[1:] == sku.replace("\r", "\n")


# --- CSV round-trip ------------------------------------------------------------------


def test_csv_round_trips_through_csv_reader() -> None:
    plan = _make_plan(
        [
            _line("REQ-R1", SupplierId.MOUSER, sku="MOUSER-SKU-1", mpn="RC0603FR-0710KL"),
        ]
    )
    result = ExportStage().run(_base_input(plan, SupplierId.MOUSER, ExportFormat.CSV), _make_ctx())
    assert result.output is not None
    text = result.output.file_bytes.decode("utf-8-sig")
    assert "\r\n" in text
    rows = list(csv.DictReader(io.StringIO(text)))
    assert rows == result.output.export.rows


# --- Determinism ----------------------------------------------------------------------


def test_export_stage_is_deterministic_across_runs() -> None:
    plan = _make_plan(
        [
            _line("REQ-R1", SupplierId.MOUSER, sku="MOUSER-SKU-1"),
            _line("REQ-C1", SupplierId.MOUSER, sku="MOUSER-SKU-2"),
        ]
    )
    ctx = _make_ctx()
    stage = ExportStage()
    inp = _base_input(plan, SupplierId.MOUSER, ExportFormat.XLSX)

    first = stage.run(inp, ctx)
    second = stage.run(inp, ctx)

    assert first.output is not None
    assert second.output is not None
    assert first.output.export.model_dump_json() == second.output.export.model_dump_json()
    assert first.output.file_bytes == second.output.file_bytes
    assert [(i.severity, i.code, i.message) for i in first.issues] == [
        (i.severity, i.code, i.message) for i in second.issues
    ]
