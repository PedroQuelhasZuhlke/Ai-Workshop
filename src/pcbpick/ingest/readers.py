"""Hardened CSV/XLSX readers: turn already-loaded file bytes into a grid of strings.

Reading `content` off disk is the runner's job (the one sanctioned I/O boundary for a
non-adapter stage); everything here is pure decoding of bytes already in memory. XLS
(legacy binary) is not implemented — no fixture exercises it yet, and an untested
reader is worse than an honest `IngestFileError`.
"""

import csv
import io

from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from pcbpick.errors import IngestFileError

MAX_FILE_SIZE_BYTES = 20 * 1024 * 1024
MAX_ROWS = 5_000
MAX_SHEETS = 20

_TEXT_ENCODINGS = ("utf-8-sig", "utf-8", "cp1252")
_CSV_DELIMITERS = ",;\t|"
_SHEET_NAME_HINTS = ("bom", "bill of materials", "parts")


def read_grid(file_name: str, content: bytes) -> list[list[str]]:
    """Dispatch on file extension to the appropriate hardened reader."""
    if len(content) > MAX_FILE_SIZE_BYTES:
        raise IngestFileError(f"{file_name} exceeds the {MAX_FILE_SIZE_BYTES}-byte ingest limit")

    suffix = file_name.rsplit(".", 1)[-1].lower() if "." in file_name else ""
    if suffix == "csv":
        return _read_csv(file_name, content)
    if suffix in ("xlsx", "xlsm"):
        return _read_xlsx(file_name, content)
    raise IngestFileError(f"unsupported BOM file type: {file_name!r}")


def _read_csv(file_name: str, content: bytes) -> list[list[str]]:
    text = _decode_text(file_name, content)

    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=_CSV_DELIMITERS)
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ","

    rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    if len(rows) > MAX_ROWS:
        raise IngestFileError(f"{file_name} exceeds the {MAX_ROWS}-row ingest limit")
    return rows


def _decode_text(file_name: str, content: bytes) -> str:
    for encoding in _TEXT_ENCODINGS:
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise IngestFileError(f"could not decode {file_name} with any supported encoding")


def _read_xlsx(file_name: str, content: bytes) -> list[list[str]]:
    # data_only=True: cached computed values, never live formula evaluation.
    workbook = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    try:
        if len(workbook.sheetnames) > MAX_SHEETS:
            raise IngestFileError(f"{file_name} exceeds the {MAX_SHEETS}-sheet ingest limit")
        sheet = _pick_sheet(workbook)
        rows: list[list[str]] = []
        for row in sheet.iter_rows(values_only=True):
            if len(rows) >= MAX_ROWS:
                raise IngestFileError(f"{file_name} exceeds the {MAX_ROWS}-row ingest limit")
            rows.append(["" if cell is None else str(cell) for cell in row])
        return rows
    finally:
        workbook.close()


def _pick_sheet(workbook: Workbook) -> Worksheet:
    if len(workbook.sheetnames) == 1:
        return workbook[workbook.sheetnames[0]]
    for name in workbook.sheetnames:
        if name.strip().lower() in _SHEET_NAME_HINTS:
            return workbook[name]
    return workbook[workbook.sheetnames[0]]
